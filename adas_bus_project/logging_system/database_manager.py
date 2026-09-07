"""SQLite persistence and tabular exports for alert events."""

from __future__ import annotations

import csv
import sqlite3
from pathlib import Path


class DatabaseManager:
    COLUMNS = ("timestamp", "track_id", "object_type", "zone", "risk_level", "event_type")

    def __init__(self, database_path: Path) -> None:
        database_path.parent.mkdir(parents=True, exist_ok=True)
        self.database_path = database_path
        with self._connect() as connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS events (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    timestamp TEXT NOT NULL,
                    track_id INTEGER,
                    object_type TEXT NOT NULL,
                    zone TEXT NOT NULL,
                    risk_level TEXT NOT NULL,
                    event_type TEXT NOT NULL
                )
                """
            )

    def _connect(self) -> sqlite3.Connection:
        return sqlite3.connect(self.database_path)

    def add_event(self, event: dict) -> None:
        values = [event.get(column) for column in self.COLUMNS]
        with self._connect() as connection:
            connection.execute(
                f"INSERT INTO events ({', '.join(self.COLUMNS)}) VALUES (?, ?, ?, ?, ?, ?)",
                values,
            )

    def recent_events(self, limit: int = 50) -> list[dict]:
        with self._connect() as connection:
            connection.row_factory = sqlite3.Row
            rows = connection.execute(
                "SELECT * FROM events ORDER BY id DESC LIMIT ?", (limit,)
            ).fetchall()
        return [dict(row) for row in rows]

    def export_csv(self, destination: Path) -> None:
        rows = self.recent_events(100000)
        with destination.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=("id", *self.COLUMNS))
            writer.writeheader()
            writer.writerows(rows)

    def export_xlsx(self, destination: Path) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise RuntimeError("Install openpyxl to export Excel reports.") from exc

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ADAS Events"
        headers = ("ID", "Timestamp", "Track ID", "Object", "Zone", "Risk", "Event")
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="16232E")
        for row in self.recent_events(100000):
            sheet.append([row.get(name) for name in ("id", *self.COLUMNS)])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(42, max(len(str(cell.value or "")) for cell in column) + 2)
            sheet.column_dimensions[column[0].column_letter].width = width
        workbook.save(destination)
